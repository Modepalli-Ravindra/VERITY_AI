import io
import os
import sys
import unittest
from pathlib import Path
from fastapi.testclient import TestClient

root_dir = Path(__file__).resolve().parent.parent
if str(root_dir) not in sys.path:
    sys.path.insert(0, str(root_dir))

from backend.app.main import app
from backend.api.auth import get_current_user
from backend.services.document_parser import (
    DocumentParser,
    MAX_FILE_SIZE_BYTES,
    MIN_CHAR_COUNT,
    MAX_DOCUMENT_CHUNKS
)
from backend.ml.text_chunking import chunk_document_bounded
import docx
import pypdf


class FakeUser:
    def __init__(self, uid: str = "test-doc-user"):
        self.id = uid
        self.email = "test@verity.ai"


class TestDocumentIntelligence(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.dependency_overrides[get_current_user] = lambda: FakeUser()
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls):
        app.dependency_overrides.clear()

    # -------------------------------------------------------------
    # 1. Valid Files
    # -------------------------------------------------------------
    def test_01_valid_txt_upload(self):
        sample_text = (
            "Artificial intelligence models are capable of generating coherent essays, "
            "synthesizing complex information, and performing natural language processing tasks. "
            "However, human writing typically exhibits nuanced vocabulary variance and irregular cadence."
        )
        file_bytes = sample_text.encode("utf-8")
        files = {"file": ("sample.txt", file_bytes, "text/plain")}
        
        response = self.client.post("/api/analyze-file", files=files, data={"detailed": "true"})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("ai_probability", data)
        self.assertIn("human_probability", data)
        self.assertIn("classification", data)
        self.assertIn("document_meta", data)
        self.assertEqual(data["document_meta"]["filename"], "sample.txt")
        self.assertEqual(data["document_meta"]["file_type"], "TXT")
        self.assertGreater(len(data.get("segments", [])), 0)

    def test_02_valid_docx_upload(self):
        doc = docx.Document()
        doc.add_paragraph(
            "Natural language generation has evolved rapidly over recent years. "
            "Large language models utilize transformer self-attention mechanisms to predict the next token. "
            "Stylometric analysis examines punctuation rhythms, stopword distributions, and vocabulary diversity."
        )
        doc.add_paragraph(
            "In forensic linguistics, human texts often feature spontaneous deviations and structural variance. "
            "Trained machine learning classifiers evaluate these signals to discern provenance."
        )
        buf = io.BytesIO()
        doc.save(buf)
        file_bytes = buf.getvalue()

        files = {"file": ("report.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        response = self.client.post("/api/analyze-file", files=files, data={"detailed": "true"})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["document_meta"]["filename"], "report.docx")
        self.assertEqual(data["document_meta"]["file_type"], "DOCX")
        self.assertGreaterEqual(data["document_meta"]["chunks_analyzed"], 1)

    def test_03_valid_pdf_upload(self):
        # Create a valid synthetic PDF with text using pypdf
        writer = pypdf.PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        
        # Test document parser directly or through endpoint with valid text bytes
        # Since pypdf PdfWriter blank page has no text, let's create a minimal PDF with actual text stream
        pdf_stream = (
            b"%PDF-1.4\n"
            b"1 0 obj <</Type /Catalog /Pages 2 0 R>> endobj\n"
            b"2 0 obj <</Type /Pages /Kids [3 0 R] /Count 1>> endobj\n"
            b"3 0 obj <</Type /Page /Parent 2 0 R /Resources <</Font <</F1 4 0 R>>>> /MediaBox [0 0 612 792] /Contents 5 0 R>> endobj\n"
            b"4 0 obj <</Type /Font /Subtype /Type1 /BaseFont /Helvetica>> endobj\n"
            b"5 0 obj <</Length 180>> stream\n"
            b"BT\n/F1 12 Tf\n72 712 Td\n(This is a valid PDF document containing sufficient sentences to evaluate artificial intelligence text detection models accurately.) Tj\nET\n"
            b"endstream\nendobj\n"
            b"xref\n0 6\n0000000000 65535 f \n0000000009 00000 n \n0000000058 00000 n \n0000000115 00000 n \n0000000244 00000 n \n0000000318 00000 n \n"
            b"trailer <</Size 6 /Root 1 0 R>>\nstartxref\n550\n%%EOF\n"
        )

        files = {"file": ("whitepaper.pdf", pdf_stream, "application/pdf")}
        response = self.client.post("/api/analyze-file", files=files, data={"detailed": "true"})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["document_meta"]["filename"], "whitepaper.pdf")
        self.assertEqual(data["document_meta"]["file_type"], "PDF")

    # -------------------------------------------------------------
    # 2. Malformed Files
    # -------------------------------------------------------------
    def test_04_malformed_pdf(self):
        malformed_bytes = b"NOT_A_VALID_PDF_HEADER_CORRUPTED_STREAM_CONTENT_XYZ"
        files = {"file": ("corrupt.pdf", malformed_bytes, "application/pdf")}
        response = self.client.post("/api/analyze-file", files=files)
        self.assertEqual(response.status_code, 400)
        self.assertIn("PDF", response.json()["detail"])

    def test_05_malformed_docx(self):
        malformed_bytes = b"CORRUPTED_NON_ZIP_BYTES_WITH_DOCX_EXTENSION"
        files = {"file": ("corrupt.docx", malformed_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        response = self.client.post("/api/analyze-file", files=files)
        self.assertEqual(response.status_code, 400)
        self.assertIn("DOCX", response.json()["detail"])

    # -------------------------------------------------------------
    # 3. Unsupported File Types
    # -------------------------------------------------------------
    def test_06_unsupported_extensions(self):
        for bad_ext in ["image.png", "script.exe", "data.csv", "table.xlsx"]:
            files = {"file": (bad_ext, b"arbitrary bytes content here", "application/octet-stream")}
            response = self.client.post("/api/analyze-file", files=files)
            self.assertEqual(response.status_code, 400)
            self.assertIn("Unsupported file type", response.json()["detail"])

    # -------------------------------------------------------------
    # 4. Oversized Files
    # -------------------------------------------------------------
    def test_07_oversized_file_rejected(self):
        oversized_bytes = b"A" * (MAX_FILE_SIZE_BYTES + 1024)
        files = {"file": ("huge.txt", oversized_bytes, "text/plain")}
        response = self.client.post("/api/analyze-file", files=files)
        self.assertEqual(response.status_code, 400)
        self.assertIn("File size exceeds", response.json()["detail"])

    # -------------------------------------------------------------
    # 5. Empty Documents
    # -------------------------------------------------------------
    def test_08_zero_byte_file(self):
        files = {"file": ("empty.txt", b"", "text/plain")}
        response = self.client.post("/api/analyze-file", files=files)
        self.assertEqual(response.status_code, 400)
        self.assertIn("empty", response.json()["detail"].lower())

    def test_09_whitespace_only_file(self):
        files = {"file": ("spaces.txt", b"   \n\n\t   \n  ", "text/plain")}
        response = self.client.post("/api/analyze-file", files=files)
        self.assertEqual(response.status_code, 400)
        self.assertIn("empty or contains no extractable text", response.json()["detail"])

    def test_10_too_short_document(self):
        short_text = "Hi there"  # Under 50 chars
        files = {"file": ("short.txt", short_text.encode("utf-8"), "text/plain")}
        response = self.client.post("/api/analyze-file", files=files)
        self.assertEqual(response.status_code, 400)
        self.assertIn("too short", response.json()["detail"])

    # -------------------------------------------------------------
    # 6. Long Documents & Bounded Chunking
    # -------------------------------------------------------------
    def test_11_bounded_chunking_unit(self):
        # Generate 100 paragraphs
        long_paragraphs = [
            f"Paragraph {i}: The deployment of advanced generative artificial intelligence models has transformed educational and research practices worldwide. Evaluating authenticity requires robust stylometric and transformer embedding inspection."
            for i in range(100)
        ]
        full_text = "\n\n".join(long_paragraphs)
        
        chunks = chunk_document_bounded(full_text, max_chunk_chars=1500, max_chunks=MAX_DOCUMENT_CHUNKS)
        self.assertLessEqual(len(chunks), MAX_DOCUMENT_CHUNKS)
        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertLessEqual(len(chunk), 1600)  # within allowable char bounds

    def test_12_long_document_end_to_end(self):
        # Generate large multi-page text (~12,000 characters)
        paragraphs = [
            "Section " + str(i) + ": Deep learning models have enabled remarkable breakthroughs in machine translation, computational linguistics, and code synthesis. In this investigation, we analyze how feature fusion combines stylometric signals with bidirectional representations to provide stable, reproducible indicators of origin."
            for i in range(40)
        ]
        doc_text = "\n\n".join(paragraphs)
        self.assertGreater(len(doc_text), 10000)

        files = {"file": ("lengthy_analysis.txt", doc_text.encode("utf-8"), "text/plain")}
        response = self.client.post("/api/analyze-file", files=files, data={"detailed": "true"})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["document_meta"]["filename"], "lengthy_analysis.txt")
        self.assertGreater(data["document_meta"]["total_words"], 1000)
        self.assertGreaterEqual(data["document_meta"]["chunks_analyzed"], 5)
        self.assertLessEqual(data["document_meta"]["chunks_analyzed"], MAX_DOCUMENT_CHUNKS)
        self.assertIn("segments", data)
        self.assertEqual(len(data["segments"]), data["document_meta"]["chunks_analyzed"])



class TestDocumentIntelligenceAuth(unittest.TestCase):
    """
    Explicit verification of JWT authentication flow on /api/analyze-file,
    /api/analyze, and /api/humanize without dependency overrides.
    """
    @classmethod
    def setUpClass(cls):
        # Ensure no dependency overrides are present
        app.dependency_overrides.clear()
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls):
        app.dependency_overrides.clear()

    @staticmethod
    def _create_jwt(sub: str = "test-user-456", exp_offset: int = 3600) -> str:
        import base64
        import json
        import time
        header = {"alg": "HS256", "typ": "JWT"}
        payload = {"sub": sub, "exp": int(time.time()) + exp_offset}
        h_b64 = base64.urlsafe_b64encode(json.dumps(header).encode()).decode().rstrip("=")
        p_b64 = base64.urlsafe_b64encode(json.dumps(payload).encode()).decode().rstrip("=")
        return f"{h_b64}.{p_b64}.signature"

    def test_auth_01_analyze_file_missing_header_returns_401(self):
        sample_text = "This is a document sample for testing unauthenticated file upload flow."
        files = {"file": ("test.txt", sample_text.encode("utf-8"), "text/plain")}
        response = self.client.post("/api/analyze-file", files=files, data={"detailed": "true"})
        self.assertEqual(response.status_code, 401)
        data = response.json()
        self.assertEqual(data.get("detail"), "Missing Authorization header")

    def test_auth_02_analyze_file_invalid_header_format_returns_401(self):
        sample_text = "This is a document sample for testing malformed authorization header."
        files = {"file": ("test.txt", sample_text.encode("utf-8"), "text/plain")}
        headers = {"Authorization": "Basic dXNlcjpwYXNz"}
        response = self.client.post("/api/analyze-file", files=files, data={"detailed": "true"}, headers=headers)
        self.assertEqual(response.status_code, 401)
        data = response.json()
        self.assertIn("Invalid Authorization header format", data.get("detail", ""))

    def test_auth_03_analyze_file_expired_token_returns_401(self):
        sample_text = "This is a document sample for testing expired token authentication."
        files = {"file": ("test.txt", sample_text.encode("utf-8"), "text/plain")}
        expired_token = self._create_jwt(exp_offset=-3600)
        headers = {"Authorization": f"Bearer {expired_token}"}
        response = self.client.post("/api/analyze-file", files=files, data={"detailed": "true"}, headers=headers)
        self.assertEqual(response.status_code, 401)
        data = response.json()
        self.assertIn("Token has expired", data.get("detail", ""))

    def test_auth_04_analyze_file_valid_token_succeeds_200(self):
        doc = docx.Document()
        doc.add_paragraph(
            "Natural language models generate text using probabilistic sequence predictions. "
            "Forensic linguistic analysis investigates punctuation, sentence rhythm, and vocabulary distributions."
        )
        doc.add_paragraph(
            "This document simulates the full guide analysis for testing authenticated document upload. "
            "Verity extracts the multi-segment text and computes human and AI probabilities."
        )
        buf = io.BytesIO()
        doc.save(buf)
        file_bytes = buf.getvalue()

        files = {"file": ("Final guide.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        valid_token = self._create_jwt(exp_offset=3600)
        headers = {"Authorization": f"Bearer {valid_token}"}
        response = self.client.post("/api/analyze-file", files=files, data={"detailed": "true"}, headers=headers)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data.get("status"), "success")
        self.assertIn("ai_probability", data)
        self.assertIn("human_probability", data)
        self.assertEqual(data.get("document_meta", {}).get("filename"), "Final guide.docx")

    def test_auth_05_analyze_text_missing_header_returns_401(self):
        response = self.client.post("/api/analyze", json={"text": "A" * 100, "detailed": True})
        self.assertEqual(response.status_code, 401)
        self.assertEqual(response.json().get("detail"), "Missing Authorization header")

    def test_auth_06_humanize_text_missing_header_returns_401(self):
        response = self.client.post("/api/humanize", json={"text": "A" * 100, "tone": "natural"})
        self.assertEqual(response.status_code, 401)
        self.assertEqual(response.json().get("detail"), "Missing Authorization header")


if __name__ == "__main__":
    unittest.main()
