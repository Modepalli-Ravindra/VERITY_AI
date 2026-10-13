import io
import os
import re
import logging
from typing import Optional, Tuple
from fastapi import HTTPException

logger = logging.getLogger("verity.document_parser")

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB limit
MIN_CHAR_COUNT = 50
MAX_DOCUMENT_CHUNKS = 50  # Bounded execution: max 50 chunks evaluated


class DocumentParser:
    """
    Secure document parser for PDF, DOCX, and TXT files.
    Enforces file-type validation, file-size limits, magic-byte checks,
    empty-document checks, and corrupt/malformed file handling.
    """

    ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

    @classmethod
    def get_extension(cls, filename: str) -> str:
        _, ext = os.path.splitext(filename or "")
        return ext.lower()

    @classmethod
    def validate_file_metadata(cls, filename: str, file_size: int):
        ext = cls.get_extension(filename)
        if ext not in cls.ALLOWED_EXTENSIONS:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported file type '{ext or 'unknown'}'. Please upload a PDF, DOCX, or TXT file."
            )

        if file_size <= 0:
            raise HTTPException(
                status_code=400,
                detail="Uploaded file is empty (0 bytes)."
            )

        if file_size > MAX_FILE_SIZE_BYTES:
            raise HTTPException(
                status_code=400,
                detail=f"File size exceeds maximum limit of {MAX_FILE_SIZE_BYTES // (1024 * 1024)}MB."
            )

    @classmethod
    def extract_text(cls, file_bytes: bytes, filename: str) -> Tuple[str, str]:
        """
        Parses file_bytes according to its validated file extension.
        Returns (extracted_text, file_type).
        """
        cls.validate_file_metadata(filename, len(file_bytes))
        ext = cls.get_extension(filename)

        extracted_text = ""
        file_type = ext.lstrip(".").upper()

        if ext == ".txt":
            extracted_text = cls._parse_txt(file_bytes)
        elif ext == ".pdf":
            extracted_text = cls._parse_pdf(file_bytes)
        elif ext == ".docx":
            extracted_text = cls._parse_docx(file_bytes)
        else:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported file type '{ext}'."
            )

        # Sanitize text
        from backend.ml.text_preprocessing import preprocess_text
        cleaned_text = preprocess_text(extracted_text)

        if not cleaned_text or not cleaned_text.strip():
            raise HTTPException(
                status_code=400,
                detail="Document is empty or contains no extractable text."
            )

        if len(cleaned_text.strip()) < MIN_CHAR_COUNT:
            raise HTTPException(
                status_code=400,
                detail=f"Document text is too short for meaningful analysis (minimum {MIN_CHAR_COUNT} characters required)."
            )

        # Enforce memory safety bound on maximum extractable text length (100,000 characters)
        if len(cleaned_text) > 100000:
            logger.info(f"Bounding extracted document text from {len(cleaned_text)} to 100,000 characters for memory safety.")
            cleaned_text = cleaned_text[:100000]

        return cleaned_text, file_type

    @classmethod
    def _parse_txt(cls, file_bytes: bytes) -> str:
        encodings_to_try = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
        for enc in encodings_to_try:
            try:
                return file_bytes.decode(enc)
            except (UnicodeDecodeError, LookupError):
                continue

        raise HTTPException(
            status_code=400,
            detail="Failed to decode text file. Ensure it is valid UTF-8 text."
        )

    @classmethod
    def _parse_pdf(cls, file_bytes: bytes) -> str:
        # Magic bytes check for PDF
        if not file_bytes.startswith(b"%PDF"):
            raise HTTPException(
                status_code=400,
                detail="Malformed or invalid PDF file (invalid file signature)."
            )

        pages_text = []

        # Attempt 1: pypdf
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            # Protect against parser resource exhaustion with high page count
            if len(reader.pages) > 200:
                raise HTTPException(
                    status_code=400,
                    detail="PDF document exceeds maximum allowed limit of 200 pages."
                )
            for page in reader.pages:
                txt = page.extract_text()
                if txt and txt.strip():
                    pages_text.append(txt.strip())
        except HTTPException:
            raise
        except Exception as pypdf_err:
            logger.warning(f"pypdf extraction error: {pypdf_err}, attempting fallback")

        # Attempt 2: pdfplumber if pypdf extracted nothing or encountered issue
        if not pages_text:
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                    if len(pdf.pages) > 200:
                        raise HTTPException(
                            status_code=400,
                            detail="PDF document exceeds maximum allowed limit of 200 pages."
                        )
                    for page in pdf.pages:
                        txt = page.extract_text()
                        if txt and txt.strip():
                            pages_text.append(txt.strip())
            except HTTPException:
                raise
            except Exception as plumber_err:
                logger.warning(f"pdfplumber extraction error: {plumber_err}")

        # If both attempts failed or caught malformed data
        if not pages_text:
            # Check if document was simply empty or truly corrupt
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                if len(reader.pages) == 0:
                    raise HTTPException(status_code=400, detail="PDF document contains 0 pages.")
            except HTTPException:
                raise
            except Exception:
                raise HTTPException(
                    status_code=400,
                    detail="Failed to parse malformed or corrupt PDF file."
                )

        return "\n\n".join(pages_text)

    @classmethod
    def _parse_docx(cls, file_bytes: bytes) -> str:
        # Magic bytes check for DOCX (standard ZIP archive signature PK\x03\x04)
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise HTTPException(
                status_code=400,
                detail="Malformed or invalid DOCX file (invalid file signature)."
            )

        # Decompression bomb and archive inspection defense
        import zipfile
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                infolist = zf.infolist()
                if len(infolist) > 1000:
                    raise HTTPException(
                        status_code=400,
                        detail="Malformed or suspicious DOCX file: excessive entry count."
                    )
                total_uncompressed = sum(info.file_size for info in infolist)
                if total_uncompressed > 50 * 1024 * 1024:  # 50 MB uncompressed threshold
                    raise HTTPException(
                        status_code=400,
                        detail="Malformed or suspicious DOCX file: uncompressed size exceeds safety threshold (50MB)."
                    )
                # Check compression ratio
                compression_ratio = total_uncompressed / max(1, len(file_bytes))
                if compression_ratio > 100:
                    raise HTTPException(
                        status_code=400,
                        detail="Decompression bomb detected in DOCX archive."
                    )
                # Check path traversal
                for info in infolist:
                    fname = info.filename
                    if fname.startswith(("/", "\\")) or ".." in fname:
                        raise HTTPException(
                            status_code=400,
                            detail="Path traversal attempt detected in DOCX archive."
                        )
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(
                status_code=400,
                detail="Failed to parse malformed or corrupt DOCX file."
            )

        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
            
            # Also extract text from tables if any
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                    if row_texts:
                        paragraphs.append(" | ".join(row_texts))

            return "\n\n".join(paragraphs)
        except HTTPException:
            raise
        except Exception as e:
            logger.warning(f"DOCX extraction error: {e}")
            raise HTTPException(
                status_code=400,
                detail="Failed to parse malformed or corrupt DOCX file."
            )
